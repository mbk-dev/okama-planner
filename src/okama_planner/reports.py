"""Neutral Excel presentation of saved forecasts, independent of the calculation path."""

from __future__ import annotations

import hashlib
import json
import math
import re
from collections.abc import Sequence
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formula import Tokenizer
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from okama_planner.api import ForecastRequest
from okama_planner.horizon import effective_horizon_years
from okama_planner.localization import terminology, translate

MONEY = '#,##0.00;(#,##0.00);"-"'
PERCENT = '0.0%;(0.0%);"-"'


@dataclass(frozen=True)
class ReportBrand:
    """Presentation only. Logo is a local image, never downloaded by this module."""

    company: str = "Example Advisory"
    contact: str = "team@example.invalid"
    color: str = "244C66"
    logo: Path | None = None


def _text(sheet: Worksheet, row: int, column: int, value: Any) -> None:
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"


def _append(sheet: Worksheet, values: Sequence[Any]) -> None:
    row = sheet.max_row + 1
    for column, value in enumerate(values, 1):
        _text(sheet, row, column, value)


def _sheet(book: Workbook, name: str, brand: ReportBrand, headers: list[Any]) -> Worksheet:
    sheet = book.create_sheet(name)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(3, len(headers)))
    _text(sheet, 1, 1, f"{brand.company} | {name}")
    sheet["A1"].font = Font(name="Arial", size=18, color="FFFFFF", bold=True)
    sheet["A1"].fill = PatternFill("solid", fgColor=brand.color)
    sheet.row_dimensions[1].height = 34
    sheet["A2"] = '=IF(\'Branding\'!B3="","",\'Branding\'!B3)'
    sheet["A3"] = "Saved forecast; edit the request and rerun Planner to change calculations."
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(3, len(headers)))
    sheet.row_dimensions[3].height = 30
    sheet.freeze_panes = "A6"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:5"
    for column, header in enumerate(headers, 1):
        _text(sheet, 5, column, header)
        cell = sheet.cell(5, column)
        cell.fill = PatternFill("solid", fgColor=brand.color)
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = 24 if column > 1 else 38
    sheet.row_dimensions[5].height = 36
    return sheet


def _validate(items: Sequence[dict[str, Any]]) -> list[ForecastRequest]:
    if not 1 <= len(items) <= 2:
        raise ValueError("Provide a baseline and optionally one comparison scenario")
    requests = []
    for item in items:
        request = ForecastRequest.model_validate(item["request"])
        result = item["result"]
        encoded = json.dumps(
            request.model_dump(mode="json"), sort_keys=True, separators=(",", ":"), allow_nan=False
        )
        digest = hashlib.sha256(encoded.encode()).hexdigest()
        if result["provenance"]["input_sha256"] != digest:
            raise ValueError("Saved result does not match its request")
        if result["schema_version"] not in {"1.0", "1.1"} or result["currency"] != request.currency:
            raise ValueError("Unsupported result schema or inconsistent currency")
        if result["portfolio_mode"] != request.portfolio_mode or (
            result["schema_version"] == "1.0" and result["portfolio_mode"] != "single"
        ):
            raise ValueError("Unsupported or inconsistent portfolio mode")
        periods = pd.period_range(
            request.plan.t0, periods=12 * effective_horizon_years(request.plan), freq="M"
        )
        months = [str(p) for p in periods]
        if list(result["ledger"]["months"]) != months:
            raise ValueError("Ledger horizon does not match the request")
        chart_months = [str(periods[0] - 1), *months]
        for series in result["charts"].values():
            if [row["month"] for row in series] != chart_months:
                raise ValueError("Chart horizon does not match the request")
        _validate_actual(request, result, chart_months)
        json.dumps(result, allow_nan=False)
        requests.append(request)
    if any(
        (r.currency, r.plan.t0, effective_horizon_years(r.plan))
        != (requests[0].currency, requests[0].plan.t0, effective_horizon_years(requests[0].plan))
        for r in requests
    ):
        raise ValueError("Comparison requires the same currency, start and horizon")
    return requests



def _validate_actual(request: ForecastRequest, result: dict[str, Any], months: list[str]) -> None:
    if result["schema_version"] != "1.1":
        return
    if request.joint_history is None or request.allocation is None:
        raise ValueError("Joint result requires a joint request")
    series = [result["actual"]["monthly_summaries"], *(s["chart"] for s in result["segments"])]
    if any([row["month"] for row in rows] != months for rows in series):
        raise ValueError("Actual or segment horizon does not match the request")

def _assumptions(sheet: Worksheet, items: Sequence[dict[str, Any]], requests: list[ForecastRequest]) -> None:
    def flatten(value: Any, prefix: str = "") -> list[tuple[str, Any]]:
        if isinstance(value, dict):
            return [
                pair for key, child in value.items() for pair in flatten(child, f"{prefix}.{key}".strip("."))
            ]
        if isinstance(value, list):
            return [
                pair for i, child in enumerate(value) for pair in flatten(child, f"{prefix}[{i}] ".strip())
            ]
        return [(prefix, value)]

    for item, request in zip(items, requests, strict=True):
        data = request.model_dump(mode="json")
        data.pop("return_samples", None)
        data.pop("joint_history", None)
        for key, value in flatten(data):
            # Preserve complete assumptions instead of silently choosing a few scenario levers.
            _append(sheet, [item["label"], key, json.dumps(value, ensure_ascii=False, allow_nan=False)])
        for key, value in flatten(item["result"]["provenance"]):
            _append(sheet, [item["label"], f"provenance.{key}", json.dumps(value, ensure_ascii=False)])
    sheet.column_dimensions["B"].width = 46
    sheet.column_dimensions["C"].width = 75


def _budget(sheet: Worksheet, result: dict[str, Any]) -> None:
    end = 5 + len(result["ledger"]["lines"])
    for month in result["ledger"]["months"]:
        row = sheet.max_row + 1
        _append(sheet, [month, None, None, None, None, result["portfolio_flow"][month]])
        for column, kinds in [
            (2, ["income"]),
            (3, ["expense"]),
            (4, ["mortgage_payment", "goal_outflow", "asset_sale", "reserve_topup"]),
        ]:
            terms = [
                f"SUMIFS('Ledger'!$D$6:$D${end},'Ledger'!$A$6:$A${end},$A{row},"
                f"'Ledger'!$B$6:$B${end},\"{kind}\")"
                for kind in kinds
            ]
            sheet.cell(row, column, "=" + "+".join(terms))
        sheet.cell(row, 5, f"=SUM(B{row}:D{row})")
    sheet["D5"] = "Other ledger flows"
    sheet["F5"] = "Portfolio flow after reserves"


def annual_cash_flow(result: dict[str, Any]) -> list[dict[str, Any]]:
    """Group supplied months into calendar years; balances use each group's last month.

    Flow kinds stay separate: buffer bookkeeping must never be added to household flows.
    The last forecast year may be partial and still has a valid final balance.
    """
    years: dict[int, dict[str, Any]] = {}
    for index, month in enumerate(result["ledger"]["months"]):
        year = int(month[:4])
        years.setdefault(year, {"year": year, "end_index": index, "flows": {}})["end_index"] = index
    for line in result["ledger"]["lines"]:
        flows = years[int(line["month"][:4])]["flows"]
        kind = line["line_kind"]
        flows[kind] = flows.get(kind, 0) + line["amount"]
    return list(years.values())


def _annual_formula(
    year: int, end: int, kinds: tuple[str, ...], source: str | None, goal_id: int | None,
) -> str:
    extra = ""
    if goal_id is not None:
        extra = f",'Ledger'!$F$6:$F${end},{goal_id}"
    elif source is not None:
        escaped = source.replace('"', '""').replace("~", "~~")
        escaped = escaped.replace("*", "~*").replace("?", "~?")
        extra = f",'Ledger'!$C$6:$C${end},\"={escaped}\""
    return "=" + "+".join(
        f"SUMIFS('Ledger'!$D$6:$D${end},'Ledger'!$A$6:$A${end},\">={year}-01\","
        f"'Ledger'!$A$6:$A${end},\"<={year}-12\",'Ledger'!$B$6:$B${end},\"{kind}\"{extra})"
        for kind in kinds
    )


def _annual_balance_rows(
    request: ForecastRequest, result: dict[str, Any], language: str,
) -> list[tuple[str, list[Any], bool]]:
    joint = result["schema_version"] == "1.1"

    def side(key: str) -> list[Any]:
        if joint:
            return [row[key]["p50"] for row in result["actual"]["monthly_summaries"]][1:]
        return list(result["ledger"][f"{key}_balance"])

    rows = [("Emergency reserve balance", side("reserve"), False),
            ("Purchase buffer", side("buffer"), False)]
    if joint:
        goals = {goal.goal_id: goal.label for goal in request.plan.goals}
        rows.extend(
            (f"{translate('Savings', language)}: {goals.get(segment['goal_id'], segment['segment_id'])}",
             [point["p50"] for point in segment["chart"]][1:], True)
            for segment in result["segments"]
        )
    elif request.plan.savings_mode == "separate":
        rows.extend(
            (f"{translate('Savings', language)}: {account['goal_label']}", list(account["balance"]), True)
            for account in result["ledger"].get("buffer_by_goal", [])
            if account["goal_label"] is not None
        )
    rows.extend([
        ("Investment portfolio", [row["p50"] for row in result["charts"]["portfolio"]][1:], False),
        ("Other assets", side("non_working"), False),
        ("Liabilities", side("liability"), False),
        ("Capital", [row["p50"] for row in result["charts"]["capital"]][1:], False),
    ])
    return rows


def _annual_goal_row(goal: Any, flow: Any) -> None:
    kinds = ("reserve_topup",) if goal.kind == "reserve_topup" else ("goal_outflow",)
    flow(goal.label, kinds, source=goal.label, literal=True, goal_id=goal.goal_id)


def _annual_income_labels(request: ForecastRequest) -> list[str]:
    """SUMIFS text criteria are case insensitive; print each matched group once."""
    labels: dict[str, str] = {}
    for item in request.plan.budget_items:
        if item.kind == "income":
            labels.setdefault(item.label.casefold(), item.label)
    return list(labels.values())


def _annual_planned_rows(request: ForecastRequest, flow: Any) -> None:
    flow("Household expenses", ("expense",))
    purchases = [goal for goal in request.plan.goals if goal.kind != "retirement_income"]
    for goal in purchases[:1]:
        _annual_goal_row(goal, flow)
    flow("Replenishing reserve", ("reserve_topup",))
    for goal in purchases[1:]:
        _annual_goal_row(goal, flow)
    for goal in request.plan.goals:
        if goal.kind == "retirement_income":
            _annual_goal_row(goal, flow)
    flow("TOTAL EXPENSES", ("expense", "goal_outflow", "reserve_topup", "mortgage_payment"))
    for label in _annual_income_labels(request):
        flow(label, ("income",), source=label, literal=True)
    flow("TOTAL INCOME", ("income", "asset_sale"))
    flow("Surplus cash", ("income", "asset_sale", "expense", "goal_outflow",
                          "reserve_topup", "mortgage_payment"))


def _annual_ages(
    sheet: Worksheet, request: ForecastRequest, annual: list[dict[str, Any]], language: str,
) -> None:
    for person in request.plan.persons:
        row = sheet.max_row + 1
        _text(sheet, row, 1, f"{person.name}: {translate('Age', language)}")
        for column, point in enumerate(annual, 2):
            sheet.cell(row, column, point["year"] - person.birth_year).number_format = "0"


def _annual_principal(sheet: Worksheet, result: dict[str, Any], annual: list[dict[str, Any]],
                      language: str) -> None:
    principal = result["ledger"].get("principal_repaid", [])
    if not principal:
        return
    row = sheet.max_row + 1
    _text(sheet, row, 1, translate("Loan principal payments", language))
    for column, point in enumerate(annual, 2):
        sheet.cell(row, column, -sum(
            value for month, value in zip(result["ledger"]["months"], principal, strict=True)
            if int(month[:4]) == point["year"]
        ))


def _annual_savings_contributions(
    request: ForecastRequest, result: dict[str, Any], language: str, flow: Any,
) -> None:
    """Show internal account deposits separately from household spending and income."""
    if result["schema_version"] != "1.0" or request.plan.savings_mode != "separate":
        return
    goals = {goal.label: goal.goal_id for goal in request.plan.goals}
    for account in result["ledger"].get("buffer_by_goal", []):
        label = account["goal_label"]
        if label is not None:
            flow(f"{translate('Annual contributions', language)}: {label}", ("buffer_in",),
                 source=label, literal=True, goal_id=goals.get(label))


def _ambiguous_reserve(kinds: tuple[str, ...], literal: bool, goal_id: int | None,
                       goal_ids: dict[int | None, int | None]) -> bool:
    return literal and kinds == ("reserve_topup",) and (goal_id is None or goal_id not in goal_ids)


def _cash_flow(
    book: Workbook, request: ForecastRequest, result: dict[str, Any], brand: ReportBrand,
    language: str,
) -> Worksheet:
    """Annual client view: planned ledger requirements and saved year-end balances."""
    annual = annual_cash_flow(result)
    sheet = _sheet(book, "Cash Flow", brand, [
        f"Cash flow ({request.currency})", *[point["year"] for point in annual],
    ])
    sheet.column_dimensions["A"].width = 48
    for column in range(2, len(annual) + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    _text(sheet, 4, 1, translate(
        "Annual flows are summed; balances are at year end (last available month). "
        "Flows are planned requirements; actual p50 balances are separate and not additive.", language,
    ))
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max(3, len(annual) + 1))
    sheet.row_dimensions[4].height = 48
    end = max(6, 5 + len(result["ledger"]["lines"]))

    goal_ids = {line.get("goal_id"): line.get("goal_id") for line in result["ledger"]["lines"]}

    def caption(label: str, literal: bool) -> int:
        row = sheet.max_row + 1
        _text(sheet, row, 1, label if literal else translate(label, language))
        return row

    def flow(label: str, kinds: tuple[str, ...], *, source: str | None = None,
             literal: bool = False, goal_id: int | None = None) -> int | None:
        if _ambiguous_reserve(kinds, literal, goal_id, goal_ids):
            return None
        row = caption(label, literal)
        goal_id = goal_ids.get(goal_id)
        for column, point in enumerate(annual, 2):
            year = point["year"]
            sheet.cell(row, column, _annual_formula(year, end, kinds, source, goal_id))
        return row

    def balance(label: str, values: list[Any], *, literal: bool = False) -> None:
        row = caption(label, literal)
        for column, point in enumerate(annual, 2):
            sheet.cell(row, column, values[point["end_index"]])

    _annual_ages(sheet, request, annual, language)
    _annual_planned_rows(request, flow)
    balance_rows = _annual_balance_rows(request, result, language)
    for label, values, literal in balance_rows[:-2]:
        balance(label, values, literal=literal)
    _annual_principal(sheet, result, annual, language)
    for label, values, literal in balance_rows[-2:]:
        balance(label, values, literal=literal)
    flow("Portfolio contributions / withdrawals", ("portfolio_flow",))
    _annual_savings_contributions(request, result, language, flow)
    return sheet


def _balance(sheet: Worksheet, result: dict[str, Any]) -> None:
    ledger = result["ledger"]
    for i, (portfolio, capital) in enumerate(
        zip(result["charts"]["portfolio"], result["charts"]["capital"], strict=True)
    ):
        row = [portfolio["month"], portfolio["p50"], capital["p50"]]
        if result["schema_version"] == "1.1":
            actual = result["actual"]["monthly_summaries"][i]
            row.extend(actual[key]["p50"] for key in ("buffer", "reserve", "non_working", "liability"))
        else:
            row.extend(
                ledger[key][i - 1] if i else None
                for key in ("buffer_balance", "reserve_balance", "non_working_balance", "liability_balance")
            )
        _append(sheet, row)
    sheet["A4"] = (
        "Actual scenario p50 balances, including opening balances; component medians are not additive."
        if result["schema_version"] == "1.1"
        else "Opening component balances are omitted; p50 series include the opening point."
    )
    sheet.merge_cells("A4:G4")
    sheet.row_dimensions[4].height = 30


def _goals(sheet: Worksheet, items: Sequence[dict[str, Any]]) -> None:
    for item in items:
        for goal in item["result"]["goals"]:
            _append(
                sheet,
                [
                    item["label"],
                    goal["label"],
                    goal["month"],
                    goal["amount_nominal"],
                    goal["p_affordable"],
                    goal["p_alive"],
                ],
            )
    sheet["A4"] = "Affordability and survival use the month before the goal; neither is full-plan success."
    sheet.merge_cells("A4:F4")


def _summary(sheet: Worksheet, result: dict[str, Any]) -> None:
    for label, value in [
        ("Full-plan success", result["metrics"]["probability_of_success"]),
        ("Terminal portfolio p50", result["metrics"]["terminal_p50"]),
        ("Portfolio mode", result["portfolio_mode"]),
        *(([("Total unmet mean", result["metrics"]["unmet_mean"])])
          if result["schema_version"] == "1.1" else
          [("Gamma", "Not available"), ("Equivalent annual alpha", "Not available")]),
    ]:
        _append(sheet, [label, value])
    sheet["B6"].number_format = PERCENT
    sheet["A12"] = "Monte Carlo percentiles describe separate distributions, not one jointly realised path."
    sheet.merge_cells("A12:D12")
    sheet.row_dimensions[12].height = 38


def _comparison(sheet: Worksheet, items: Sequence[dict[str, Any]]) -> None:
    baseline = items[0]["result"]
    other = items[1]["result"] if len(items) == 2 else None
    for key, label in [
        ("probability_of_success", "Full-plan success"),
        ("terminal_p50", "Terminal portfolio p50"),
    ]:
        row = sheet.max_row + 1
        values = [label, baseline["metrics"][key], other["metrics"][key] if other else "Not supplied"]
        _append(sheet, [*values, None if other else "Not supplied"])
        if other:
            sheet.cell(row, 4, f"=C{row}-B{row}")
    if all(item["result"]["schema_version"] == "1.0" for item in items):
        for label in ["Gamma", "Equivalent annual alpha"]:
            _append(sheet, [label, "Not available", "Not available", "Not available"])
    else:
        _joint_comparison(sheet, items)
    for column in ("B", "C", "D"):
        sheet[f"{column}6"].number_format = PERCENT
    sheet["A4"] = (
        "Differences are descriptive; changed goals or cash flows prevent a like-for-like investment ranking."
    )
    sheet.merge_cells("A4:D4")
    sheet.row_dimensions[4].height = 36



def _joint_goals(sheet: Worksheet, items: Sequence[dict[str, Any]]) -> None:
    for item in items:
        for goal in item["result"]["goals"]:
            joint = item["result"]["schema_version"] == "1.1"
            _append(sheet, [
                item["label"], goal["label"], goal["month"], goal["amount_nominal"],
                goal["p_funded"] if joint else "Legacy: see affordability in saved JSON",
                goal.get("p_full_stream"), goal.get("funded_mean"), goal.get("unmet_mean"),
                goal.get("funding_basis", "legacy_before_goal_affordability"),
            ])
    sheet["A4"] = "Actual funding covers every required goal event; pension success covers the full stream."
    sheet.merge_cells("A4:I4")


def _joint_comparison(sheet: Worksheet, items: Sequence[dict[str, Any]]) -> None:
    left = items[0]["result"]
    right = items[1]["result"] if len(items) == 2 else None

    def difference(label: str, a: Any, b: Any, probability: bool = False) -> None:
        row = sheet.max_row + 1
        _append(sheet, [label, a, b, None])
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            sheet.cell(row, 4, f"=C{row}-B{row}")
        else:
            _text(sheet, row, 4, "Not comparable")
        if probability:
            for column in (2, 3, 4):
                sheet.cell(row, column).number_format = PERCENT

    difference("Total unmet mean", left["metrics"].get("unmet_mean", "Legacy: unavailable"),
               right["metrics"].get("unmet_mean", "Legacy: unavailable") if right else "Not supplied")
    other_goals = {g["goal_id"]: g for g in right["goals"]} if right else {}
    for goal in left["goals"]:
        other = other_goals.get(goal["goal_id"], {})
        for key, label, probability in [
            ("p_funded", "fully funded probability", True),
            ("p_full_stream", "full-stream probability", True),
            ("funded_mean", "funded mean", False),
            ("unmet_mean", "unmet mean", False),
        ]:
            if key == "p_full_stream" and goal.get(key) is None and other.get(key) is None:
                continue
            difference(f"{goal['label']}: {label}", goal.get(key, "Unavailable"),
                       other.get(key, "Unavailable"), probability)
    for label, key in [("Portfolio mode", "portfolio_mode"), ("Shared history hash", "history_sha256"),
                       ("Shared scenario rows hash", "scenario_rows_sha256")]:
        if key == "portfolio_mode":
            values = [r[key] if r else "Not supplied" for r in (left, right)]
        else:
            values = [r["provenance"].get(key, "Legacy: unavailable") if r else "Not supplied"
                      for r in (left, right)]
        _append(sheet, [label, *values, "Same" if values[0] == values[1] else "Different"])


def _joint_details(book: Workbook, items: Sequence[dict[str, Any]], brand: ReportBrand) -> None:
    segments = _sheet(book, "Segments", brand, [
        "Scenario", "Segment", "Goal ID", "Opening amount", "Full-funding probability", "Completion",
    ])
    allocation = _sheet(book, "Allocation", brand, ["Scenario", "Policy / schedule", "Explicit value"])
    balances = _sheet(book, "Segment balances", brand, ["Scenario", "Segment", "Month", "p10", "p50", "p90"])
    events = _sheet(book, "Funding events", brand, [
        "Scenario", "Month", "Kind", "Label", "Goal ID", "Segment", "Required mean", "Funded mean",
        "Unmet mean", "Fully funded probability", "Funding sources (mean)",
    ])
    transfers = _sheet(book, "Transfers", brand, [
        "Scenario", "Mode", "Month", "Source", "Destination", "Reason", "Amount mean",
    ])
    segments["A4"] = "Single mode segment balances are a display attribution within one pooled portfolio."
    segments.merge_cells("A4:F4")
    allocation.column_dimensions["C"].width = 95
    for item in items:
        result = item["result"]
        if result["schema_version"] != "1.1":
            continue
        label = item["label"]
        for key, value in result["allocation"].items():
            if key != "segments":
                _append(allocation, [label, key, json.dumps(value, ensure_ascii=False)])
        for segment in result["segments"]:
            _append(segments, [label, segment["segment_id"], segment["goal_id"], segment["opening_amount"],
                               segment["probability_of_full_funding"], json.dumps(segment["completion"])])
            segments.cell(segments.max_row, 5).number_format = PERCENT
            for key in ("strategy", "active_strategy"):
                _append(allocation, [label, f"{segment['segment_id']}.{key}", json.dumps(segment[key])])
            for point in segment["chart"]:
                _append(balances, [label, segment["segment_id"], point["month"],
                                   point["p10"], point["p50"], point["p90"]])
        for event in result["actual"]["event_funding"]:
            _append(events, [label, event["month"], event["kind"], event["label"], event["goal_id"],
                             event["segment_id"], event["required_mean"], event["funded_mean"],
                             event["unmet_mean"], event["p_funded"], json.dumps(event["funding_sources"])])
            events.cell(events.max_row, 10).number_format = PERCENT
        for transfer in result["actual"]["transfers"]:
            _append(transfers, [label, result["portfolio_mode"], transfer["month"], transfer["source"],
                                transfer["destination"], transfer["reason"], transfer["amount_mean"]])

def _style_cell(cell: Any) -> None:
    formula = cell.data_type == "f"
    color = "008000" if formula and "!" in cell.value else "000000" if formula else "0000FF"
    cell.font = Font(name="Arial", size=11, color=color)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if isinstance(cell.value, str) and not formula:
        cell.data_type = "s"
    if isinstance(cell.value, (float, int)) or formula and cell.row >= 6:
        if cell.number_format == "General":
            cell.number_format = MONEY
    if not formula and isinstance(cell.value, (float, int)):
        cell.comment = Comment(
            "Source: supplied saved forecast; see Assumptions for request hash and provenance.",
            "okama Planner",
        )


def _finish(book: Workbook) -> None:
    for sheet in book:
        for row in sheet:
            for cell in row:
                if sheet.title != "Branding" and cell.row in (1, 5):
                    continue
                _style_cell(cell)
            if row[0].row > 5:
                lines = max(
                    math.ceil(
                        (12 if cell.data_type == "f" else len(str(cell.value or "")))
                        / max(10, sheet.column_dimensions[get_column_letter(cell.column)].width)
                    )
                    for cell in row
                )
                sheet.row_dimensions[row[0].row].height = min(120, max(18, 16 * lines))
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions
    for coordinate in ("B2", "B3"):
        book["Branding"][coordinate].fill = PatternFill("solid", fgColor="FFF2CC")


def _print_layout(book: Workbook, summary: Worksheet, unit: str, label: str) -> None:
    book.move_sheet(summary, offset=-1)
    branding = book["Branding"]
    branding.page_setup.orientation = "landscape"
    branding.page_setup.paperSize = branding.PAPERSIZE_A4
    branding.page_setup.fitToWidth = 1
    branding.page_setup.fitToHeight = 1
    branding.sheet_properties.pageSetUpPr.fitToPage = True
    branding.print_area = "A1:F7"
    branding.column_dimensions["A"].width = 26
    for sheet in book:
        sheet.oddFooter.right.text = f"Nominal {unit} | {label}"
        sheet.oddFooter.right.font = "Arial"
    for name in ("Budget", "Ledger"):
        basis = " | Planned requirements, not actual payments/balances" if "Funding events" in book else ""
        book[name]["A4"] = f"Baseline: {label} | Nominal {unit}{basis}"
        book[name].merge_cells(start_row=4, start_column=1, end_row=4, end_column=book[name].max_column)


def _template_cell(sheet: str, row: int, column: int) -> bool:
    """Select template cells by role, never by the content of user labels."""
    if sheet == "Branding":
        return column == 1 or (row in (1, 5) and column == 2) or (row == 5 and column == 4)
    if row <= 5:
        return not (sheet == "Comparison" and row == 5 and column in (2, 3))
    if sheet in {"Assumptions", "Ledger", "Balance", "Budget", "Allocation", "Segments",
                 "Segment balances", "Funding events", "Transfers", "Cash Flow", "Current amounts"}:
        return False
    if sheet == "Goals":
        return column in (5, 6, 7, 8)
    return True


def _localized_text(value: str, language: str, unit: str, goal_metric: bool = False) -> str:
    texts = terminology(language)
    if value in texts:
        return texts[value]
    if value.replace(unit, "{unit}") in texts:
        return translate(value.replace(unit, "{unit}"), language, unit=unit)
    if goal_metric:
        for suffix in ("fully funded probability", "full-stream probability", "funded mean", "unmet mean"):
            if value.endswith(f": {suffix}"):
                return value[:-len(suffix)] + translate(suffix, language)
    return value


def _localized_formula(value: str, names: dict[str, str]) -> str:
    """Rename reference tokens only; text criteria contain opaque user labels."""
    tokens = Tokenizer(value).items
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        for old, new in names.items():
            prefix = f"'{old}'!"
            if token.value.startswith(prefix):
                escaped = new.replace("'", "''")
                token.value = f"'{escaped}'!" + token.value[len(prefix):]
                break
    return "=" + "".join(token.value for token in tokens)


def _localize_cells(
    sheet: Worksheet, names: dict[str, str], language: str, unit: str, brand: ReportBrand,
) -> None:
    for row in sheet:
        for cell in row:
            if cell.data_type == "f":
                cell.value = _localized_formula(cell.value, names)
            elif isinstance(cell.value, str) and _template_cell(sheet.title, cell.row, cell.column):
                value = (f"{brand.company} | {names[sheet.title]}"
                         if sheet.title != "Branding" and cell.coordinate == "A1"
                         else _localized_text(cell.value, language, unit,
                                              sheet.title == "Comparison" and cell.column == 1))
                _text(sheet, cell.row, cell.column, value)
            if cell.comment:
                cell.comment = Comment(translate(cell.comment.text, language), cell.comment.author)


def _localized_layout(sheet: Worksheet, language: str, resize: bool = True) -> None:
    for row in sheet:
        if language == "zh":
            for cell in row:
                font = copy(cell.font)
                font.name = "Microsoft YaHei"
                cell.font = font
        if resize and row[0].row >= 5:
            lines = max(math.ceil((12 if cell.data_type == "f" else len(str(cell.value or ""))) /
                                  max(10, sheet.column_dimensions[get_column_letter(cell.column)].width))
                        for cell in row)
            old_height = sheet.row_dimensions[row[0].row].height or 18
            sheet.row_dimensions[row[0].row].height = min(120, max(old_height, 16 * lines))


def _localize(
    book: Workbook, language: str, unit: str, brand: ReportBrand, scenarios: Sequence[dict[str, Any]],
) -> None:
    names = {sheet.title: translate(sheet.title, language) for sheet in book}
    for sheet in book:
        original_name = sheet.title
        _localize_cells(sheet, names, language, unit, brand)
        _localized_layout(sheet, language)
        if original_name == "Comparison" and len(scenarios) == 1:
            _text(sheet, 5, 3, translate("No comparison", language))
        if original_name in {"Ledger", "Budget"}:
            basis = (" · " + translate("Planned requirements, not actual payments/balances", language)
                     if "Funding events" in names else "")
            _text(sheet, 4, 1, translate("Baseline: {label} · Nominal {unit}", language,
                                       label=str(scenarios[0]["label"]), unit=unit) + basis)
        sheet.oddFooter.right.text = translate("Nominal {unit} · {label}", language,
                                               unit=unit, label=str(scenarios[0]["label"]))
        sheet.title = names[original_name]
    book.properties.title = translate("Neutral financial plan", language)


def _style_chart_sheet(sheet: Worksheet, language: str) -> None:
    for row in sheet:
        for cell in row:
            if cell.row != 1:
                _style_cell(cell)
    _localized_layout(sheet, language, resize=False)


def _chart_sheets(
    book: Workbook, images: dict[str, str | Path], language: str, unit: str, brand: ReportBrand,
    goals: list[dict[str, Any]],
) -> None:
    names = {"portfolio": "Portfolio chart", "portfolio_log": "Portfolio log chart",
             "capital": "Capital chart", "capital_log": "Capital log chart"}
    if set(images) - names.keys():
        raise ValueError(f"Unknown chart image keys: {sorted(set(images) - names.keys())}")
    for key, name in names.items():
        if key not in images:
            continue
        sheet = _sheet(book, translate(name, language), brand, [])
        sheet.freeze_panes = None
        sheet.unmerge_cells("A1:C1")
        sheet.merge_cells("A1:P1")
        sheet.unmerge_cells("A3:C3")
        contact = f"'{translate('Branding', language)}'!B3"
        sheet["A2"] = f'=IF({contact}="","",{contact})'
        _text(sheet, 3, 1, translate(
            "Nominal {unit}; inflation indexation follows the supplied inputs. "
            "Values are not expressed in constant purchasing power.", language, unit=unit,
        ))
        sheet.merge_cells("A3:P3")
        sheet.row_dimensions[3].height = 48
        if key.endswith("_log"):
            _text(sheet, 4, 1, translate(
                "Zero and negative values are omitted on the logarithmic scale; affected bands have gaps.",
                language,
            ))
            sheet.merge_cells("A4:P4")
            sheet.row_dimensions[4].height = 30
        image = Image(images[key])
        image.width, image.height = 1200, 720
        sheet.add_image(image, "A6")
        for column in range(1, 17):
            sheet.column_dimensions[get_column_letter(column)].width = 11
        for row in range(6, 43):
            sheet.row_dimensions[row].height = 15
        for index, goal in enumerate(goals, 1):
            row = 43 + index
            _text(sheet, row, 1, f"{index} — {goal['label']} ({goal['month']})")
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=16)
            sheet.row_dimensions[row].height = 24
        sheet.print_area = f"A1:P{43 + len(goals)}"
        sheet.page_setup.fitToHeight = 1
        _style_chart_sheet(sheet, language)


def _current_amounts(book: Workbook, request: ForecastRequest, brand: ReportBrand, language: str) -> None:
    sheet = _sheet(book, "Current amounts", brand, [
        "Category", "Label", f"Current amount ({request.currency})", "Period / basis", "Annual rate",
    ])
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["D"].width = 30
    plan = request.plan
    records = [
        ("Asset", asset.label, asset.amount, "Opening balance", asset.growth_rate, False)
        for asset in plan.assets
    ] + [
        ("Liability", loan.label, loan.principal, "Opening principal", loan.annual_rate, False)
        for loan in plan.liabilities
    ] + [
        ("Goal", goal.label, goal.amount_pv,
         f"{goal.pv_year}: " + translate(
             "Expense share" if goal.amount_basis == "expense_share" else "Present value", language),
         goal.indexation_rate, goal.amount_basis == "expense_share")
        for goal in plan.goals
    ] + [
        ({"income": "Income", "expense": "Expenses"}.get(item.kind, item.kind),
         item.label, item.monthly_amount, "Monthly", item.indexation_rate, False)
        for item in plan.budget_items
    ]
    for category, label, amount, basis, rate, ratio in records:
        _append(sheet, [translate(category, language), label, amount, translate(basis, language), rate])
        sheet.cell(sheet.max_row, 5).number_format = PERCENT
        if ratio:
            sheet.cell(sheet.max_row, 3).number_format = PERCENT


def export_report(
    scenarios: Sequence[dict[str, Any]],
    path: str | Path,
    *,
    brand: ReportBrand | None = None,
    language: str = "en",
    chart_images: dict[str, str | Path] | None = None,
) -> Path:
    """Export one saved forecast and optionally a comparison. Calculations are never rerun.

    Each scenario has ``label``, ``request`` and ``result``. Only identical units/horizons
    can be compared. Branding and strings are presentation inputs, never Excel formulas.
    """
    terminology(language)
    requests = _validate(scenarios)
    brand = brand or ReportBrand()
    if not re.fullmatch(r"[0-9a-fA-F]{6}", brand.color):
        raise ValueError("Brand color must be six hexadecimal digits")
    book = Workbook()
    book.remove(book.active)
    branding = book.create_sheet("Branding")
    for row, (label, value) in enumerate(
        [
            ("Branding", "Presentation settings"),
            ("Company", brand.company),
            ("Contact", brand.contact),
            ("Color", brand.color),
            ("Editable", "B2/B3: text only; rerun exporter for logo/color changes"),
        ],
        1,
    ):
        _text(branding, row, 1, label)
        _text(branding, row, 2, value)
    branding.column_dimensions["B"].width = 70
    if brand.logo:
        image = Image(brand.logo)
        image.height, image.width = 48, 144
        branding.add_image(image, "D2")
        branding["D5"] = "Local example logo"
    _current_amounts(book, requests[0], brand, language)
    result = scenarios[0]["result"]
    unit = requests[0].currency
    summary = _sheet(book, "Summary", brand, ["Indicator", f"Baseline ({unit})", "", ""])
    _summary(summary, result)
    _budget(
        _sheet(
            book,
            "Budget",
            brand,
            ["Month", "Income", "Expenses", "Other flows", "Net ledger flow", "Portfolio flow"],
        ),
        result,
    )
    _balance(
        _sheet(
            book,
            "Balance",
            brand,
            [
                "Month",
                "Portfolio p50",
                "Net capital p50",
                "Purchase buffer",
                "Reserve account",
                "Non-working assets",
                "Debt",
            ],
        ),
        result,
    )
    joint = any(item["result"]["schema_version"] == "1.1" for item in scenarios)
    (_joint_goals if joint else _goals)(
        _sheet(
            book,
            "Goals",
            brand,
            [
                "Scenario",
                "Goal",
                "Month",
                f"Nominal amount ({unit})",
                "Affordable probability",
                "Alive probability",
            ] if not joint else [
                "Scenario", "Goal", "First month", f"First required amount ({unit})",
                "Fully funded probability", "Full-stream probability", "Funded mean (whole goal)",
                "Unmet mean (whole goal)", "Funding basis",
            ],
        ),
        scenarios,
    )
    _assumptions(
        _sheet(book, "Assumptions", brand, ["Scenario", "Parameter", "Supplied value"]), scenarios, requests
    )
    _comparison(
        _sheet(
            book,
            "Comparison",
            brand,
            [
                "Indicator",
                str(scenarios[0]["label"]),
                str(scenarios[1]["label"]) if len(scenarios) == 2 else "No comparison",
                "Difference",
            ],
        ),
        scenarios,
    )
    raw = _sheet(book, "Ledger", brand, [
        "Month", "Kind", "Label", f"Amount ({unit})", "Resolved rate", "Goal ID",
    ])
    for line in result["ledger"]["lines"]:
        _append(raw, [line["month"], line["line_kind"], line["label"], line["amount"],
                      line["resolved_rate"], line.get("goal_id")])
    cash_flow = _cash_flow(book, requests[0], result, brand, language)
    book.move_sheet(cash_flow, offset=2 - book.sheetnames.index("Cash Flow"))
    if joint:
        _joint_details(book, scenarios, brand)
    instructions = _sheet(book, "Instructions", brand, ["Topic", "Instruction", ""])
    for row in [
        [
            "Edit",
            "Edit Branding B2/B3 for contact text. Rerun the exporter to change "
            "company headings, colors and local logo.",
        ],
        [
            "Calculations",
            "Request/result are snapshots. Change the full request and rerun forecast before exporting; "
            "editing workbook numbers does not rerun Monte Carlo.",
        ],
        [
            "Assumptions",
            "Every scenario has its own full plan and provenance. Return histories are identified by hashes; "
            "keep the original request/result JSON files.",
        ],
        [
            "Units",
            f"All monetary amounts are nominal {unit}; inflation indexation follows the supplied inputs. "
            "Rates and probabilities are fractions. "
            "Budget expenses and other outflows are negative.",
        ],
        [
            "Limits",
            "Legacy single or joint single/per_goal investments; fixed-rate savings are separate reserves. "
            "No jurisdictional tax, FX, gamma/alpha or white label legal declaration is inferred.",
        ],
        [
            "Country rules",
            "Supply supported budgets and cash flows explicitly. "
            "Fees/taxes absent from inputs are not modelled. "
            "Formatting does not change financial assumptions.",
        ],
        [
            "Recalculate",
            "Recalculate Excel summary formulas before using cached values. "
            "Compare numeric series separately; "
            "percentile components are not additive.",
        ],
    ]:
        _append(instructions, row)
    instructions.column_dimensions["B"].width = 95
    _finish(book)
    _print_layout(book, summary, unit, str(scenarios[0]["label"]))
    for row in range(6, book["Goals"].max_row + 1):
        for column in (5, 6):
            book["Goals"].cell(row, column).number_format = PERCENT
    for row in range(6, raw.max_row + 1):
        raw.cell(row, 5).number_format = PERCENT
    book.active = book.sheetnames.index("Summary")
    book.properties.creator = "okama Planner"
    book.properties.title = "Neutral financial plan"
    _localize(book, language, unit, brand, scenarios)
    _chart_sheets(book, chart_images or {}, language, unit, brand, result["goals"])
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)
    return destination
