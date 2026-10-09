"""Saved multi-currency forecasts with native group views and explicit exchange traces."""

from __future__ import annotations

import json
import re
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image

from okama_planner.api import _digest
from okama_planner.charts import _chart_data
from okama_planner.horizon import effective_horizon_years
from okama_planner.ledger.budget import budget_lines
from okama_planner.ledger.calendar import year_index
from okama_planner.ledger.types import LedgerLine
from okama_planner.localization import translate
from okama_planner.multicurrency import MulticurrencyRequest
from okama_planner.reports import (
    PERCENT,
    ReportBrand,
    _append,
    _chart_sheets,
    _finish,
    _localize,
    _sheet,
    _text,
)


def _requests(items: Sequence[dict[str, Any]]) -> list[MulticurrencyRequest]:
    if not 1 <= len(items) <= 2:
        raise ValueError("Provide a baseline and optionally one comparison scenario")
    requests = []
    for item in items:
        request = MulticurrencyRequest.model_validate(item["request"])
        result = item["result"]
        if (
            result.get("schema_version") != "2.0"
            or result["currency"] != request.currency
            or result["provenance"]["input_sha256"] != _digest(request.model_dump(mode="json"))
        ):
            raise ValueError("Saved multi-currency result does not match its request")
        json.dumps(result, allow_nan=False)
        _chart_data(result)
        expected = pd.period_range(
            request.household.t0,
            periods=12
            * effective_horizon_years(
                request.household,
            ),
            freq="M",
        )
        months = [str(expected[0] - 1), *[str(m) for m in expected]]
        _validate_group_charts(request, result, months)
        requests.append(request)
    identities = {(r.currency, r.household.t0, effective_horizon_years(r.household)) for r in requests}
    if len(identities) != 1:
        raise ValueError("Comparison requires the same currency, start and horizon")
    return requests


def _validate_group_charts(request: MulticurrencyRequest, result: dict[str, Any], months: list[str]) -> None:
    groups = result["currency_groups"]
    if [(g["group_id"], g["currency"]) for g in groups] != [
        (g.group_id, g.request.currency) for g in request.groups
    ]:
        raise ValueError("Native result groups do not match the request")
    for supplied, group in zip(request.groups, groups, strict=True):
        native = group["result"]
        if native["currency"] != supplied.request.currency:
            raise ValueError("Native result currency does not match its request")
        if native["provenance"]["input_sha256"] != _digest(supplied.request.model_dump(mode="json")):
            raise ValueError("Native result input hash does not match its request")
    for selected in [result, *[g["result"] for g in groups]]:
        _chart_data(selected)
        if any([p["month"] for p in series] != months for series in selected["charts"].values()):
            raise ValueError("Multi-currency chart horizon does not match the request")
        _validate_actual_calendar(selected, months)
    if "household_ledger" in result and list(result["household_ledger"]["months"]) != months[1:]:
        raise ValueError("Household ledger calendar does not match the request")


def _validate_actual_calendar(selected: dict[str, Any], months: list[str]) -> None:
    actual = selected["actual"]
    if [row["month"] for row in actual["monthly_summaries"]] != months:
        raise ValueError("Actual monthly calendar does not match the request")
    if any(row["month"] not in months[1:] for row in actual["event_funding"]):
        raise ValueError("Actual event calendar does not match the request")
    if any(row["month"] not in months[1:] for row in actual.get("fx_transfers", [])):
        raise ValueError("FX transfer calendar does not match the request")
    if "ledger" in selected and list(selected["ledger"]["months"]) != months[1:]:
        raise ValueError("Native ledger calendar does not match the request")


def _budget(
    book: Workbook,
    request: MulticurrencyRequest,
    brand: ReportBrand,
    result: dict[str, Any],
) -> None:
    plan = request.household
    retirement = year_index(plan.t0, plan.retirement_year)
    lines = budget_lines(
        plan.budget_items, plan.t0, 12 * effective_horizon_years(plan), retirement, plan.rates
    )
    if "household_ledger" in result:
        lines = [
            LedgerLine(**row)
            for row in result["household_ledger"]["lines"]
            if row["line_kind"] in {"income", "expense"}
        ]
    sheet = _sheet(book, "Budget", brand, ["Month", "Kind", "Label", f"Amount ({request.currency})"])
    years = sorted({int(line.month[:4]) for line in lines})
    if not years:
        years = sorted(
            {
                int(str(m)[:4])
                for m in pd.period_range(
                    plan.t0,
                    periods=12 * effective_horizon_years(plan),
                    freq="M",
                )
            }
        )
    cash = _sheet(book, "Cash Flow", brand, [f"Cash flow ({request.currency})", *years])
    for line in lines:
        _append(sheet, [line.month, line.line_kind, line.label, line.amount])
    for caption, kind in [("TOTAL INCOME", "income"), ("Household expenses", "expense")]:
        _append(
            cash,
            [
                caption,
                *[
                    sum(
                        row[3]
                        for row in sheet.iter_rows(min_row=6, values_only=True)
                        if row[1] == kind and row[0].startswith(str(year))
                    )
                    for year in years
                ],
            ],
        )
    _text(cash, 8, 1, "Free cash flow")
    for col in range(2, len(years) + 2):
        from openpyxl.utils import get_column_letter

        letter = get_column_letter(col)
        cash.cell(8, col, f"=SUM({letter}6:{letter}7)")


def _native_flow(
    book: Workbook,
    group: dict[str, Any],
    overall: dict[str, Any],
    brand: ReportBrand,
    language: str,
) -> None:
    result, currency = group["result"], group["currency"]
    points = result["charts"]["portfolio"][1:]
    years = sorted({int(p["month"][:4]) for p in points})
    sheet = _sheet(book, f"Cash Flow {currency}", brand, [f"Cash flow ({currency})", *years])
    summaries = result["actual"]["monthly_summaries"]
    for caption, field, sign in [
        ("Currency contributions", "contribution_mean", 1),
        ("Household support", "household_withdrawal_mean", -1),
    ]:
        _append(
            sheet,
            [
                caption,
                *[
                    sign * sum(row.get(field, 0) for row in summaries if row["month"].startswith(str(year)))
                    for year in years
                ],
            ],
        )
    loans = [e for e in result["actual"]["event_funding"] if e["kind"] == "mortgage_payment"]
    if loans:
        _append(
            sheet,
            [
                "Mortgage payments",
                *[
                    -sum(e["funded_mean"] for e in loans if e["month"].startswith(str(year)))
                    for year in years
                ],
            ],
        )
    financing = [
        t
        for t in result["actual"].get("fx_transfers", [])
        if t["direction"] == "contribution" and t["reason"] != "surplus"
    ]
    if financing:
        _append(
            sheet,
            [
                "Loan financing",
                *[
                    sum(t["native_amount_mean"] for t in financing if t["month"].startswith(str(year)))
                    for year in years
                ],
            ],
        )
    for goal in result["goals"]:
        events = [e for e in result["actual"]["event_funding"] if e["goal_id"] == goal["goal_id"]]
        _append(
            sheet,
            [
                goal["label"],
                *[
                    -sum(e["funded_mean"] for e in events if e["month"].startswith(str(year)))
                    for year in years
                ],
            ],
        )
    for caption, key in [("Investment portfolio", "portfolio"), ("Capital", "capital")]:
        _append(
            sheet,
            [
                caption,
                *[
                    [p["p50"] for p in result["charts"][key] if p["month"].startswith(str(year))][-1]
                    for year in years
                ],
            ],
        )

    goal_rows = {8 + bool(loans) + bool(financing) + i for i in range(len(result["goals"]))}
    for row in range(6, sheet.max_row + 1):
        if row not in goal_rows:
            _text(sheet, row, 1, translate(str(sheet.cell(row, 1).value), language))


def _groups_and_goals(
    book: Workbook, request: MulticurrencyRequest, result: dict[str, Any], brand: ReportBrand, language: str
) -> None:
    groups = _sheet(
        book,
        "Currency groups",
        brand,
        ["Group", "Currency", "Portfolio mode", "Probability of success", "Portfolio p50"],
    )
    inputs = _sheet(
        book, "Current amounts", brand, ["Group", "Currency", "Category", "Label", "Amount", "Period / basis"]
    )
    strategies = _sheet(book, "Allocation", brand, ["Group", "Currency", "Strategy", "Supplied value"])
    for group, saved in zip(request.groups, result["currency_groups"], strict=True):
        native = group.request
        _append(
            groups,
            [
                group.group_id,
                native.currency,
                native.portfolio_mode,
                saved["result"]["metrics"]["probability_of_success"],
                saved["result"]["charts"]["portfolio"][-1]["p50"],
            ],
        )
        groups.cell(groups.max_row, 4).number_format = PERCENT
        for asset in native.plan.assets:
            _append(inputs, [group.group_id, native.currency, "Asset", asset.label, asset.amount])
        for liability in native.plan.liabilities:
            _append(
                inputs, [group.group_id, native.currency, "Liability", liability.label, liability.principal]
            )
        for goal in native.plan.goals:
            _append(
                inputs,
                [
                    group.group_id,
                    native.currency,
                    "Goal",
                    goal.label,
                    goal.amount_pv,
                    "Expense share" if goal.amount_basis == "expense_share" else "Present value",
                ],
            )
            if goal.amount_basis == "expense_share":
                inputs.cell(inputs.max_row, 5).number_format = PERCENT
        for key, value in native.allocation.model_dump(mode="json").items():
            _append(strategies, [group.group_id, native.currency, key, json.dumps(value, ensure_ascii=False)])
        _native_flow(book, saved, result, brand, language)
    goals = _sheet(
        book,
        "Goals",
        brand,
        [
            "Group",
            "Currency",
            "Goal",
            "Month",
            "Nominal amount",
            "Fully funded probability",
            "Funded mean (whole goal)",
            "Unmet mean (whole goal)",
        ],
    )
    for goal in result["goals"]:
        _append(
            goals,
            [
                goal["group_id"],
                goal["currency"],
                goal["label"],
                goal["month"],
                goal["amount_nominal"],
                goal["p_funded"],
                goal["funded_mean"],
                goal["unmet_mean"],
            ],
        )
        goals.cell(goals.max_row, 6).number_format = PERCENT


def _traces(
    book: Workbook, request: MulticurrencyRequest, result: dict[str, Any], brand: ReportBrand
) -> None:
    fx = _sheet(
        book,
        "FX transfers",
        brand,
        [
            "Month",
            "Group",
            "Source currency",
            "Destination currency",
            "Source amount",
            "Destination amount",
            "Exchange rate",
            f"Fee ({request.currency})",
            "Reason",
        ],
    )
    for transfer in result["actual"]["fx_transfers"]:
        _append(
            fx,
            [
                transfer.get(key)
                for key in (
                    "month",
                    "group_id",
                    "source_currency",
                    "destination_currency",
                    "source_amount_mean",
                    "destination_amount_mean",
                    "rate_mean",
                    "fee_base_mean",
                    "reason",
                )
            ],
        )
    events = _sheet(
        book,
        "Funding events",
        brand,
        [
            "Month",
            "Group",
            "Currency",
            "Label",
            "Kind",
            "Required mean",
            "Funded mean",
            "Unmet mean",
            f"Unmet mean ({request.currency})",
        ],
    )
    for event in result["actual"]["event_funding"]:
        _append(
            events,
            [
                event.get(key)
                for key in (
                    "month",
                    "group_id",
                    "currency",
                    "label",
                    "kind",
                    "required_mean",
                    "funded_mean",
                    "unmet_mean",
                    "unmet_base_mean",
                )
            ],
        )
    assumptions = _sheet(book, "Assumptions", brand, ["Parameter", "Supplied value", ""])
    data = request.model_dump(mode="json")
    for key, value in data.items():
        if key == "groups":
            value = [
                {**g, "request": {k: v for k, v in g["request"].items() if k != "joint_history"}}
                for g in value
            ]
        elif key == "fx":
            value = {k: v for k, v in value.items() if k != "monthly_returns"}
        encoded = json.dumps(value, ensure_ascii=False)
        # Chunk large JSON for Excel's per-cell limit; histories remain in request snapshots.
        for offset in range(0, len(encoded), 30000):
            _append(
                assumptions, [key if not offset else f"{key} ({offset})", encoded[offset : offset + 30000]]
            )
    _append(assumptions, ["provenance", json.dumps(result["provenance"], ensure_ascii=False)])


def export_multicurrency_report(
    scenarios: Sequence[dict[str, Any]],
    path: str | Path,
    *,
    brand: ReportBrand | None = None,
    language: str = "en",
    chart_images: dict[str, str | Path] | None = None,
) -> Path:
    requests = _requests(scenarios)
    request, result = requests[0], scenarios[0]["result"]
    brand = brand or ReportBrand()
    if not re.fullmatch(r"[0-9a-fA-F]{6}", brand.color):
        raise ValueError("Brand color must be six hexadecimal digits")
    book = Workbook()
    book.remove(book.active)
    _branding(book, brand)
    summary = _sheet(book, "Summary", brand, ["Indicator", f"Baseline ({request.currency})", "", ""])
    for caption, value in [
        ("Probability of success", result["metrics"]["probability_of_success"]),
        ("Portfolio p50", result["charts"]["portfolio"][-1]["p50"]),
        ("Net capital p50", result["charts"]["capital"][-1]["p50"]),
    ]:
        _append(summary, [caption, value])
    summary["B6"].number_format = PERCENT
    _budget(book, request, brand, result)
    _groups_and_goals(book, request, result, brand, language)
    _traces(book, request, result, brand)
    if len(scenarios) == 2:
        comparison = _sheet(book, "Comparison", brand, ["Indicator", "Baseline", "Variant", "Difference"])
        for metric in ("probability_of_success", "terminal_p50"):
            left, right = (s["result"]["metrics"][metric] for s in scenarios)
            _append(comparison, [metric, left, right, right - left])
    instructions = _sheet(book, "Instructions", brand, ["Topic", "Instruction", ""])
    _append(
        instructions, ["Units", "Group amounts stay in native currencies; totals are converted pathwise."]
    )
    _append(
        instructions,
        ["Calculations", "FX transfers are internal; fees reduce capital. Snapshots retain histories."],
    )
    _finish(book)
    _localize(book, language, request.currency, brand, scenarios)
    for group in result["currency_groups"]:
        old = f"Cash Flow {group['currency']}"
        sheet = book[old]
        sheet.title = f"{translate('Cash Flow', language)} {group['currency']}"[:31]
        _text(
            sheet,
            3,
            1,
            translate(
                "Nominal {unit}; inflation indexation follows the supplied inputs. "
                "Values are not expressed in constant purchasing power.",
                language,
                unit=group["currency"],
            ),
        )
        _text(sheet, 5, 1, translate("Cash flow ({unit})", language, unit=group["currency"]))
        sheet.oddFooter.right.text = translate(
            "Nominal {unit} · {label}", language, unit=group["currency"], label=str(scenarios[0]["label"])
        )
    _group_chart_sheets(book, chart_images or {}, language, request.currency, brand, result)
    book.active = book.sheetnames.index(translate("Summary", language))
    book.properties.creator = "okama Planner"
    book.properties.title = translate("Neutral financial plan", language)
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return target


def _group_chart_sheets(
    book: Workbook,
    images: dict[str, str | Path],
    language: str,
    currency: str,
    brand: ReportBrand,
    result: dict[str, Any],
) -> None:
    keys = {"portfolio", "portfolio_log", "capital", "capital_log"}
    allowed = keys | {f"{g['group_id']}/{key}" for g in result["currency_groups"] for key in keys}
    if set(images) - allowed:
        raise ValueError(f"Unknown chart image keys: {sorted(set(images) - allowed)}")
    _chart_sheets(
        book, {k: v for k, v in images.items() if k in keys}, language, currency, brand, result["goals"]
    )
    for group in result["currency_groups"]:
        selected = {
            key: images[f"{group['group_id']}/{key}"]
            for key in keys
            if f"{group['group_id']}/{key}" in images
        }
        _chart_sheets(
            book,
            selected,
            language,
            group["currency"],
            brand,
            group["result"]["goals"],
            title_suffix=f" {group['currency']}",
        )


def _branding(book: Workbook, brand: ReportBrand) -> None:
    branding = book.create_sheet("Branding")
    for row, values in enumerate(
        [
            ("Branding", "Presentation settings"),
            ("Company", brand.company),
            ("Contact", brand.contact),
            ("Color", brand.color),
        ],
        1,
    ):
        for column, value in enumerate(values, 1):
            _text(branding, row, column, value)
    if brand.logo:
        branding.add_image(Image(brand.logo), "D2")
