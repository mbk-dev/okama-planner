"""Check reports against independently saved public calculations."""

import json
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

EXAMPLES = Path(__file__).parents[1] / "examples"


def scenario(name: str = "baseline") -> dict:
    return {
        "label": name.title(),
        "request": json.loads((EXAMPLES / f"{name}-request.json").read_text()),
        "result": json.loads((EXAMPLES / f"{name}-result.json").read_text()),
    }


def test_report_preserves_money_and_distinguishes_capital(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    target = tmp_path / "report.xlsx"
    export_report([scenario(), scenario("deferred")], target)
    book = load_workbook(target)
    assert {"Summary", "Budget", "Balance", "Goals", "Assumptions", "Comparison"} <= set(book.sheetnames)
    assert book["Summary"]["B6"].value == 0.974
    assert book["Summary"]["B7"].value == pytest.approx(25291.057126370935)
    assert book["Comparison"]["D7"].value == "=C7-B7"
    for cell in ("B6", "C6", "D6"):
        assert book["Budget"][cell].data_type == "f"
        assert "Ledger" in book["Budget"][cell].value
    assert book["Budget"]["E6"].value == "=SUM(B6:D6)"
    assert book["Budget"].row_dimensions[6].height <= 24
    assert book["Balance"]["B6"].value == 75000
    assert book["Balance"]["C6"].value == 75000
    assert "Portfolio p50" in book["Balance"]["B5"].value
    assert "Net capital p50" in book["Balance"]["C5"].value
    assert book["Balance"]["E5"].value == "Reserve account"
    assert book["Summary"]["B9"].value == "Not available"
    assert book["Summary"]["B10"].value == "Not available"


def test_brand_changes_presentation_without_changing_results(tmp_path: Path) -> None:
    from okama_planner.reports import ReportBrand, export_report

    first, second = tmp_path / "one.xlsx", tmp_path / "two.xlsx"
    export_report([scenario()], first)
    export_report(
        [scenario()],
        second,
        brand=ReportBrand(company="=Example Advisory", contact="team@example.invalid", color="703080"),
    )
    a, b = load_workbook(first), load_workbook(second)
    assert b["Branding"]["B2"].value == "=Example Advisory"
    assert b["Branding"]["B2"].data_type == "s"
    assert b["Summary"]["A1"].fill.fgColor.rgb.endswith("703080")
    assert [a["Summary"][f"B{r}"].value for r in range(6, 11)] == [
        b["Summary"][f"B{r}"].value for r in range(6, 11)
    ]


@pytest.mark.parametrize("mutation", ["request", "currency", "schema", "months"])
def test_report_rejects_inconsistent_saved_calculation(tmp_path: Path, mutation: str) -> None:
    from okama_planner.reports import export_report

    item = deepcopy(scenario())
    if mutation == "request":
        item["request"]["seed"] += 1
    elif mutation == "currency":
        item["result"]["currency"] = "EUR"
    elif mutation == "schema":
        item["result"]["schema_version"] = "9.0"
    else:
        item["result"]["charts"]["capital"].pop()
    with pytest.raises(ValueError):
        export_report([item], tmp_path / "bad.xlsx")
    assert not (tmp_path / "bad.xlsx").exists()


def test_comparison_discloses_changed_inputs_and_rejects_different_units(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    export_report([scenario(), scenario("deferred")], tmp_path / "compare.xlsx")
    book = load_workbook(tmp_path / "compare.xlsx")
    rows = list(book["Assumptions"].values)
    assert any("target_year" in str(row) and "2029" in str(row) for row in rows)
    item = scenario("deferred")
    item["request"]["currency"] = "EUR"
    for asset in item["request"]["plan"]["assets"]:
        asset["currency"] = "EUR"
    from okama_planner import forecast

    item["result"] = forecast(item["request"])
    with pytest.raises(ValueError, match="Comparison requires"):
        export_report([scenario(), item], tmp_path / "different.xlsx")


def test_logo_is_embedded_and_color_validated(tmp_path: Path) -> None:
    from PIL import Image

    from okama_planner.reports import ReportBrand, export_report

    logo = tmp_path / "logo.png"
    Image.new("RGB", (144, 48), "purple").save(logo)
    export_report([scenario()], tmp_path / "logo.xlsx", brand=ReportBrand(logo=logo))
    import zipfile

    with zipfile.ZipFile(tmp_path / "logo.xlsx") as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())
    with pytest.raises(ValueError, match="color"):
        export_report([scenario()], tmp_path / "bad.xlsx", brand=ReportBrand(color="invalid"))


def test_scenario_labels_are_text_in_every_sheet(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    item = scenario()
    item["label"] = "=1+1"
    export_report([item], tmp_path / "labels.xlsx")
    book = load_workbook(tmp_path / "labels.xlsx")
    for name in ("Goals", "Assumptions"):
        assert book[name]["A6"].value == "=1+1"
        assert book[name]["A6"].data_type == "s"


def joint_scenario(mode: str = "per_goal", opening: float = 100) -> dict:
    from okama_planner import forecast, with_portfolio_mode
    from test_allocation import allocation, request

    raw = request(opening)
    raw["plan"]["budget_items"] = []
    raw["plan"]["goals"] = [
        {"goal_id": 1, "label": "Home", "kind": "lump", "amount_pv": 80,
         "pv_year": 2026, "target_year": 2026, "becomes_asset": True},
        {"goal_id": 2, "label": "Pension", "kind": "retirement_income",
         "amount_pv": 20, "pv_year": 2026},
    ]
    raw["allocation"] = allocation(opening, (1, 2))
    raw["allocation"]["segments"][1]["completion"] = {
        "action": "transfer_to", "destination": "g2"
    }
    if mode == "single":
        raw = with_portfolio_mode(raw, portfolio_mode=mode).model_dump(mode="json")
    return {"label": mode, "request": raw, "result": forecast(raw)}


def test_joint_report_uses_actual_funding_and_side_balances(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    item = joint_scenario(opening=0)
    export_report([item], tmp_path / "actual.xlsx")
    book = load_workbook(tmp_path / "actual.xlsx")
    assert book["Balance"]["F7"].value == 0  # Planned property is 80; failed purchase owns none.
    assert book["Balance"]["D6"].value == 0  # Actual opening side balance is available.
    assert "actual" in book["Balance"]["A4"].value.lower()
    assert "planned" in book["Budget"]["A4"].value.lower()
    assert "planned" in book["Ledger"]["A4"].value.lower()
    assert book["Goals"]["E6"].value == 0
    assert book["Goals"]["G6"].value == 0
    assert book["Goals"]["H6"].value == 80
    assert book["Goals"]["I6"].value == "actual_event"
    assert book["Goals"]["F7"].value == 0
    assert book["Goals"]["H7"].value == 240
    assert book["Goals"]["I7"].value == "full_stream"
    assert not any(row[0] in {"Gamma", "Equivalent annual alpha"} for row in book["Summary"].values)


def test_joint_comparison_discloses_segments_events_and_differences(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    export_report([joint_scenario("single"), joint_scenario()], tmp_path / "modes.xlsx")
    book = load_workbook(tmp_path / "modes.xlsx")
    assert {"Segments", "Allocation", "Segment balances", "Funding events", "Transfers"} <= set(
        book.sheetnames
    )
    pension = list(book["Goals"].values)[6]
    assert pension[4:9] == (0, 0, 20, 220, "full_stream")
    assert any("transfer_to" in str(row) and "g2" in str(row) for row in book["Segments"].values)
    assert any("active_strategy" in str(row) and '"A"' in str(row) for row in book["Allocation"].values)
    assert any("surplus_weights" in str(row) for row in book["Allocation"].values)
    assert any(row[3:6] == ("household", "g1", "shortfall") for row in book["Transfers"].values)
    assert any(row[0] == "Total unmet mean" and row[3].startswith("=")
               for row in book["Comparison"].values)
    assert any(row[0] == "Pension: full-stream probability" for row in book["Comparison"].values)
    assert any(row[6:10] == (80, 80, 0, 1) for row in book["Funding events"].values)


@pytest.mark.parametrize("series", ["actual", "segment"])
def test_joint_report_rejects_incomplete_actual_horizon(tmp_path: Path, series: str) -> None:
    from okama_planner.reports import export_report

    item = joint_scenario()
    if series == "actual":
        item["result"]["actual"]["monthly_summaries"].pop()
    else:
        item["result"]["segments"][0]["chart"].pop()
    with pytest.raises(ValueError, match="horizon"):
        export_report([item], tmp_path / "bad.xlsx")
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize("language,summary,budget,ledger", [
    ("en", "Summary", "Budget", "Ledger"),
    ("ru", "Итоги", "Бюджет", "Операции"),
    ("zh", "摘要", "预算", "流水账"),
    ("de", "Übersicht", "Budget", "Buchungen"),
    ("es", "Resumen", "Presupuesto", "Movimientos"),
])
def test_localized_reports_preserve_labels_formulas_and_money(
    tmp_path: Path, language: str, summary: str, budget: str, ledger: str,
) -> None:
    from okama_planner.reports import export_report

    item = scenario()
    item["label"] = "Summary"  # A user label coinciding with a template term must survive.
    target = export_report([item], tmp_path / f"{language}.xlsx", language=language)
    book = load_workbook(target)
    assert book[summary]["B7"].value == pytest.approx(item["result"]["metrics"]["terminal_p50"])
    assert f"'{ledger}'!" in book[budget]["B6"].value
    assert all(sheet.freeze_panes in (None, "A6") for sheet in book)
    from okama_planner.localization import translate

    assert book[translate("Assumptions", language)]["A6"].value == "Summary"
    assert len(set(book.sheetnames)) == len(book.sheetnames)
    assert all(len(name) <= 31 for name in book.sheetnames)
    assert not any(cell.value == "License" for sheet in book for row in sheet for cell in row)
    if language != "en":
        comparison = {"ru": "Сравнение", "zh": "比较", "de": "Vergleich", "es": "Comparación"}[language]
        assert book[comparison]["C5"].value != "No comparison"
        assert book[summary]["A5"].value != "Indicator"
        assert book[summary]["A3"].value != (
            "Saved forecast; edit the request and rerun Planner to change calculations."
        )


def test_report_embeds_four_chart_images_without_freeze_panes(tmp_path: Path) -> None:
    from PIL import Image

    from okama_planner.reports import export_report

    image = tmp_path / "chart.png"
    Image.new("RGB", (1200, 720), "white").save(image)
    keys = ("portfolio", "portfolio_log", "capital", "capital_log")
    target = export_report([scenario()], tmp_path / "charts.xlsx",
                           chart_images=dict.fromkeys(keys, image))
    book = load_workbook(target)
    charts = [sheet for sheet in book if sheet._images]
    assert len(charts) == 4
    assert all(sheet.freeze_panes is None for sheet in charts)
    assert all("inflation" in sheet["A3"].value for sheet in charts)
    assert "Zero" in charts[1]["A4"].value


def test_empty_contact_is_blank_formula_not_zero(tmp_path: Path) -> None:
    from okama_planner.reports import ReportBrand, export_report

    book = load_workbook(export_report([scenario()], tmp_path / "blank.xlsx",
                                     brand=ReportBrand(contact="")))
    assert book["Summary"]["A2"].value == '=IF(\'Branding\'!B3="","",\'Branding\'!B3)'


def test_unsupported_report_language_fails_before_writing(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    with pytest.raises(ValueError, match="language"):
        export_report([scenario()], tmp_path / "bad.xlsx", language="fr")
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize("language", ["en", "ru", "zh", "de", "es"])
def test_localized_company_headings_are_literal_text(tmp_path: Path, language: str) -> None:
    from okama_planner.reports import ReportBrand, export_report

    book = load_workbook(export_report([scenario()], tmp_path / "headings.xlsx", language=language,
                                      brand=ReportBrand(company="=1+1")))
    assert all(sheet["A1"].data_type == "s" for sheet in book)
    assert all(sheet["A1"].value.startswith("=1+1 | ") for sheet in book if sheet.title != {
        "en": "Branding", "ru": "Оформление", "zh": "品牌设置", "de": "Gestaltung", "es": "Marca",
    }[language])


def test_annual_cash_flow_sums_flows_and_selects_last_available_balance(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    item = scenario()
    book = load_workbook(export_report([item], tmp_path / "annual.xlsx"))
    assert "Cash Flow" in book
    sheet = book["Cash Flow"]
    assert [sheet.cell(5, c).value for c in range(2, 8)] == [2026, 2027, 2028, 2029, 2030, 2031]
    rows = {sheet.cell(r, 1).value: r for r in range(6, sheet.max_row + 1)}
    assert "SUMIFS" in sheet.cell(rows["Household expenses"], 2).value
    assert '"expense"' in sheet.cell(rows["Household expenses"], 2).value
    assert sheet.cell(rows["Investment portfolio"], 2).value == (
        item["result"]["charts"]["portfolio"][12]["p50"]
    )
    assert sheet.cell(rows["Capital"], 2).value == item["result"]["charts"]["capital"][12]["p50"]
    assert sheet.cell(rows["Emergency reserve balance"], 2).value == (
        item["result"]["ledger"]["reserve_balance"][11]
    )
    assert sheet.freeze_panes == "A6"
    assert sheet.column_dimensions["B"].width == 18


def test_annual_cash_flow_partial_year_and_negative_flows() -> None:
    from okama_planner.reports import annual_cash_flow

    result = {"ledger": {"months": ["2026-11", "2026-12", "2027-01"], "lines": [
        {"month": "2026-11", "line_kind": "expense", "amount": -30},
        {"month": "2026-12", "line_kind": "expense", "amount": -40},
        {"month": "2027-01", "line_kind": "expense", "amount": -50},
        {"month": "2026-12", "line_kind": "buffer_in", "amount": 999},
    ]}}
    assert annual_cash_flow(result) == [
        {"year": 2026, "end_index": 1, "flows": {"expense": -70, "buffer_in": 999}},
        {"year": 2027, "end_index": 2, "flows": {"expense": -50}},
    ]


@pytest.mark.parametrize("language,currency", [
    ("en", "EUR"), ("ru", "USD"), ("zh", "RUB"), ("de", "CNY"), ("es", "USD"),
])
def test_report_currency_comes_from_request_independently_of_language(
    tmp_path: Path, language: str, currency: str,
) -> None:
    from okama_planner import forecast
    from okama_planner.localization import translate
    from okama_planner.reports import export_report

    item = scenario()
    item["request"]["currency"] = currency
    for asset in item["request"]["plan"]["assets"]:
        asset["currency"] = currency
    item["result"] = forecast(item["request"])
    book = load_workbook(export_report([item], tmp_path / "currency.xlsx", language=language))
    cash_flow = book[translate("Cash Flow", language)]
    assert currency in cash_flow["A5"].value
    assert f"'{translate('Ledger', language)}'!" in cash_flow["B6"].value
    assert currency in book[translate("Current amounts", language)]["C5"].value
    assert currency in book[translate("Summary", language)]["B5"].value
    assert currency in book[translate("Goals", language)]["D5"].value
    assert currency in book[translate("Ledger", language)]["D5"].value
    assert all(currency in sheet.oddFooter.right.text for sheet in book)


def test_annual_goal_labels_remain_literal_and_use_ids(tmp_path: Path) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report

    item = scenario()
    label = '=SUM(A1:A2)*?"'
    item["request"]["plan"]["goals"][0]["label"] = label
    item["result"] = forecast(item["request"])
    book = load_workbook(export_report([item], tmp_path / "names.xlsx"))
    cell = next(row[0] for row in book["Cash Flow"] if row[0].value == label)
    assert cell.data_type == "s"
    assert "'Ledger'!$F$6:" in book["Cash Flow"].cell(cell.row, 2).value


def test_annual_legacy_goal_without_ids_uses_literal_label_filter(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    item = scenario()
    for line in item["result"]["ledger"]["lines"]:
        line.pop("goal_id", None)
    book = load_workbook(export_report([item], tmp_path / "old.xlsx"))
    assert '"=Synthetic purchase"' in book["Cash Flow"]["D7"].value
    assert "'Ledger'!$F$6:" not in book["Cash Flow"]["D7"].value


def test_annual_reserve_goals_use_reserve_lines_and_income_labels_do_not_duplicate(tmp_path: Path) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report

    item = scenario()
    item["request"]["plan"]["goals"].append({
        "goal_id": 4, "label": "Safety target", "kind": "reserve_topup", "amount_pv": 1000,
        "pv_year": 2026, "target_year": 2027,
    })
    item["request"]["plan"]["budget_items"].append({
        "kind": "income", "label": "Synthetic household income", "monthly_amount": 100,
    })
    item["result"] = forecast(item["request"])
    book = load_workbook(export_report([item], tmp_path / "reserve.xlsx"))
    sheet = book["Cash Flow"]
    rows = [row[0].value for row in sheet]
    assert rows.count("Synthetic household income") == 1
    cell = next(row[0] for row in sheet if row[0].value == "Safety target")
    assert '"reserve_topup"' in sheet.cell(cell.row, 2).value


@pytest.mark.parametrize("label", ["=Salary", ">Pay", "<Income", "*Literal?~"])
def test_annual_income_criteria_force_literal_equality(tmp_path: Path, label: str) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report

    item = scenario()
    item["request"]["plan"]["budget_items"][0]["label"] = label
    item["result"] = forecast(item["request"])
    book = load_workbook(export_report([item], tmp_path / "literal.xlsx"))
    sheet = book["Cash Flow"]
    cell = next(row[0] for row in sheet if row[0].value == label)
    escaped = label.replace("~", "~~").replace("*", "~*").replace("?", "~?")
    assert f'"={escaped}"' in sheet.cell(cell.row, 2).value


@pytest.mark.parametrize("language", ["en", "ru", "zh", "de", "es"])
def test_current_amounts_present_native_request_amounts_with_dynamic_units(
    tmp_path: Path, language: str,
) -> None:
    from okama_planner.localization import translate
    from okama_planner.reports import export_report

    item = scenario()
    book = load_workbook(export_report([item], tmp_path / "inputs.xlsx", language=language))
    sheet = book[translate("Current amounts", language)]
    assert "USD" in sheet["C5"].value
    assert sheet["C6"].value == item["request"]["plan"]["assets"][0]["amount"]
    assert any(row[2].value == 20000 for row in sheet)


def test_legacy_reserve_goals_without_ids_have_only_aggregate_row(tmp_path: Path) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report

    item = scenario()
    item["request"]["plan"]["goals"].extend([
        {"goal_id": 4, "label": "Reserve A", "kind": "reserve_topup", "amount_pv": 1000,
         "pv_year": 2026, "target_year": 2027},
        {"goal_id": 5, "label": "Reserve B", "kind": "reserve_topup", "amount_pv": 2000,
         "pv_year": 2026, "target_year": 2028},
    ])
    item["result"] = forecast(item["request"])
    for line in item["result"]["ledger"]["lines"]:
        line.pop("goal_id", None)
    book = load_workbook(export_report([item], tmp_path / "legacy-reserves.xlsx"))
    captions = [row[0].value for row in book["Cash Flow"]]
    assert "Reserve A" not in captions
    assert "Reserve B" not in captions
    assert captions.count("Replenishing reserve") == 1


@pytest.mark.parametrize("label", ["Pay 'Ledger'!", "Pay 'Branding'!"])
def test_localized_formula_sheet_references_preserve_literal_user_criteria(
    tmp_path: Path, label: str,
) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report

    item = scenario()
    item["request"]["plan"]["budget_items"][0]["label"] = label
    item["result"] = forecast(item["request"])
    book = load_workbook(export_report([item], tmp_path / "criteria.xlsx", language="ru"))
    sheet = book["Денежные потоки"]
    cell = next(row[0] for row in sheet if row[0].value == label)
    formula = sheet.cell(cell.row, 2).value
    assert f'"={label}"' in formula
    assert "'Операции'!$D$6:" in formula


def test_annual_income_labels_share_excel_case_insensitive_groups(tmp_path: Path) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report

    item = scenario()
    item["request"]["plan"]["budget_items"][0]["label"] = "Salary"
    item["request"]["plan"]["budget_items"].append({
        "kind": "income", "label": "salary", "monthly_amount": 100,
    })
    item["result"] = forecast(item["request"])
    book = load_workbook(export_report([item], tmp_path / "salary.xlsx"))
    captions = [row[0].value for row in book["Cash Flow"]]
    assert captions.count("Salary") == 1
    assert "salary" not in captions
    assert any(row[1].value == "salary" for row in book["Current amounts"])
