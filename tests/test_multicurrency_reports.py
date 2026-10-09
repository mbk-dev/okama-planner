"""Presentation keeps native values separate from base-currency totals."""

from pathlib import Path

from openpyxl import load_workbook

from okama_planner.api import _digest
from test_multicurrency import multicurrency_request


def saved_multicurrency() -> dict:
    from okama_planner.multicurrency import MulticurrencyRequest

    request = MulticurrencyRequest.model_validate(multicurrency_request())
    from okama_planner.ledger.calendar import month_key

    months = [month_key(request.household.t0, n) for n in range(-1, 24)]

    def chart(value: float) -> list[dict]:
        return [{"month": month, **{f"p{p}": value for p in (10, 25, 50, 75, 90)}} for month in months]

    groups = [
        {
            "group_id": g.group_id,
            "currency": g.request.currency,
            "result": {
                "schema_version": "1.1",
                "currency": g.request.currency,
                "portfolio_mode": g.request.portfolio_mode,
                "charts": {"portfolio": chart(100), "capital": chart(100)},
                "metrics": {"probability_of_success": 1.0},
                "goals": [],
                "actual": {"monthly_summaries": [{"month": m} for m in months], "event_funding": []},
                "provenance": {"input_sha256": _digest(g.request.model_dump(mode="json"))},
            },
        }
        for g in request.groups
    ]
    result = {
        "schema_version": "2.0",
        "currency": "RUB",
        "currency_groups": groups,
        "charts": {"portfolio": chart(10100), "capital": chart(10100)},
        "goals": [],
        "metrics": {"probability_of_success": 1.0, "terminal_p50": 10100},
        "actual": {
            "event_funding": [],
            "fx_transfers": [],
            "monthly_summaries": [{"month": m} for m in months],
        },
        "provenance": {"input_sha256": _digest(request.model_dump(mode="json"))},
    }
    return {"label": "Family", "request": request.model_dump(mode="json"), "result": result}


def test_multicurrency_html_exports_overall_and_each_native_group(tmp_path: Path) -> None:
    from okama_planner.charts import export_charts

    paths = export_charts(saved_multicurrency()["result"], tmp_path)
    assert len(paths) == 3
    assert (tmp_path / "dollar" / "forecast.html").is_file()
    assert '"currency":"USD"' in (tmp_path / "dollar" / "forecast.html").read_text()
    assert '"currency":"RUB"' in (tmp_path / "forecast.html").read_text()


def test_multicurrency_excel_has_native_years_and_base_summary(tmp_path: Path) -> None:
    from okama_planner.reports import export_report

    export_report([saved_multicurrency()], tmp_path / "plan.xlsx")
    book = load_workbook(tmp_path / "plan.xlsx")
    assert {"Summary", "Currency groups", "Cash Flow USD", "FX transfers", "Goals"} <= set(book.sheetnames)
    native = book["Cash Flow USD"]
    assert native["B5"].value == 2026
    assert native["A5"].value.endswith("USD)")
    assert native.freeze_panes == "A6"
    rows = {native.cell(r, 1).value: r for r in range(6, native.max_row + 1)}
    assert native.cell(rows["Investment portfolio"], 2).value == 100
    assert book["Summary"]["B7"].value == 10100


def test_native_chart_images_have_native_currency_captions(tmp_path: Path) -> None:
    from PIL import Image
    from okama_planner.reports import export_report

    image = tmp_path / "chart.png"
    Image.new("RGB", (50, 30), "white").save(image)
    export_report(
        [saved_multicurrency()],
        tmp_path / "native.xlsx",
        chart_images={"portfolio": image, "dollar/portfolio": image},
    )
    book = load_workbook(tmp_path / "native.xlsx")
    assert len(book["Portfolio chart USD"]._images) == 1
    assert "USD" in book["Portfolio chart USD"]["A3"].value


def test_chart_group_ids_cannot_escape_output_directory(tmp_path: Path) -> None:
    import pytest
    from okama_planner.charts import export_charts

    result = saved_multicurrency()["result"]
    result["currency_groups"][0]["group_id"] = "../outside"
    with pytest.raises(ValueError, match="safe"):
        export_charts(result, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_native_cashflow_uses_actual_contributions_and_household_withdrawals(tmp_path: Path) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report

    item = saved_multicurrency()
    item["result"] = forecast(item["request"])
    export_report([item], tmp_path / "actual.xlsx")
    book = load_workbook(tmp_path / "actual.xlsx")
    sheet = book["Cash Flow USD"]
    rows = {sheet.cell(r, 1).value: r for r in range(6, sheet.max_row + 1)}
    assert sheet.cell(rows["Currency contributions"], 2).value == 30
    assert "USD" in sheet.oddFooter.right.text


def test_localization_preserves_native_goal_labels(tmp_path: Path) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report
    from test_multicurrency_engine import add_goal, monthly_budget

    raw = multicurrency_request()
    monthly_budget(raw, 0, 0)
    group = add_goal(raw, 1, 1, 20)
    group["plan"]["goals"][0]["label"] = "Capital"
    item = {"label": "Family", "request": raw, "result": forecast(raw)}
    export_report([item], tmp_path / "ru.xlsx", language="ru")
    book = load_workbook(tmp_path / "ru.xlsx")
    assert "Capital" in [c.value for c in book["Денежные потоки USD"]["A"]]


def test_native_cashflow_exposes_loan_financing(tmp_path: Path) -> None:
    from okama_planner import forecast
    from okama_planner.reports import export_report
    from test_multicurrency_engine import monthly_budget

    raw = multicurrency_request()
    monthly_budget(raw, 1000, 0)
    raw["contribution_schedule"][0]["weights"] = {"ruble": 1, "dollar": 0}
    raw["conversion_fee_rate"] = 0.1
    raw["groups"][1]["request"]["plan"]["liabilities"] = [
        {
            "label": "Loan",
            "principal": 9,
            "annual_rate": 0,
            "monthly_payment": 9,
            "term_months": 1,
            "start_month": "2026-01",
        }
    ]
    item = {"label": "Family", "request": raw, "result": forecast(raw)}
    export_report([item], tmp_path / "loan.xlsx")
    sheet = load_workbook(tmp_path / "loan.xlsx")["Cash Flow USD"]
    rows = {sheet.cell(r, 1).value: r for r in range(6, sheet.max_row + 1)}
    assert sheet.cell(rows["Loan financing"], 2).value == 9


def test_report_rejects_native_actuals_outside_horizon(tmp_path: Path) -> None:
    import pytest
    from okama_planner import forecast
    from okama_planner.reports import export_report

    raw = multicurrency_request()
    result = forecast(raw)
    result["currency_groups"][1]["result"]["actual"]["monthly_summaries"][1]["month"] = "2035-01"
    with pytest.raises(ValueError, match="calendar"):
        export_report([{"label": "Family", "request": raw, "result": result}], tmp_path / "bad.xlsx")


def test_report_rejects_nested_currency_mismatch(tmp_path: Path) -> None:
    import pytest
    from okama_planner.reports import export_report

    item = saved_multicurrency()
    item["result"]["currency_groups"][1]["result"]["currency"] = "EUR"
    with pytest.raises(ValueError, match="currency"):
        export_report([item], tmp_path / "bad.xlsx")
