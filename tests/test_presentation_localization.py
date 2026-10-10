"""Presentation localization preserves machine values and rejects incomplete catalogs."""

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from test_multicurrency_reports import saved_multicurrency


def test_catalog_rejects_missing_translation_duplicates_and_placeholders() -> None:
    from okama_planner.localization import validate_terminology

    good = {
        "key": "test",
        "en": "Amount {unit}",
        "ru": "Сумма {unit}",
        "zh": "金额 {unit}",
        "de": "Betrag {unit}",
        "es": "Importe {unit}",
    }
    validate_terminology([good])
    for bad in ({**good, "ru": ""}, {**good, "es": "Importe {currency}"}):
        with pytest.raises(ValueError):
            validate_terminology([bad])
    with pytest.raises(ValueError):
        validate_terminology([good, good])


def test_packaged_catalog_is_complete() -> None:
    from okama_planner.localization import validate_terminology
    from importlib.resources import files

    with files("okama_planner").joinpath("terminology.csv").open(encoding="utf-8") as source:
        validate_terminology(list(csv.DictReader(source)))


@pytest.mark.parametrize(
    ("language", "number", "date"),
    [
        ("en", "1,234.50", "10/09/2026"),
        ("ru", "1\u00a0234,50", "09.10.2026"),
        ("de", "1.234,50", "09.10.2026"),
        ("es", "1.234,50", "09/10/2026"),
        ("zh", "1,234.50", "2026年10月9日"),
    ],
)
def test_locale_formats_are_explicit(language: str, number: str, date: str) -> None:
    from okama_planner.localization import format_date, format_number

    assert format_number(1234.5, language) == number
    assert format_date("2026-10-09", language) == date
    assert format_date("2026-10", language) != ""


def test_client_presentation_preserves_input_and_translates_enums() -> None:
    from okama_planner.localization import client_presentation

    record = {"code": "c-0001", "full_name": "Goals", "sex": "female", "ips_sent_at": "2026-10-09"}
    shown = {item["key"]: item for item in client_presentation(record, "ru")}
    assert shown["full_name"]["label"] == "Полное имя"
    assert shown["full_name"]["value"] == "Goals"
    assert shown["sex"]["value"] == "Женский"
    assert shown["ips_sent_at"]["value"] == "09.10.2026"
    assert record["sex"] == "female"


def test_validation_messages_do_not_include_private_inputs() -> None:
    from okama_planner.localization import localized_validation_errors, localized_error
    from okama_planner.storage.validation import ClientDetails

    try:
        ClientDetails(full_name="private name", sex="private invalid value")
    except ValidationError as error:
        rows = localized_validation_errors(error, "de")
        assert rows[0]["loc"] == ["sex"]
        assert "private" not in str(rows)
        assert "private" not in localized_error(error, "de")
        assert rows[0]["msg"] != error.errors()[0]["msg"]


@pytest.mark.parametrize("language", ["ru", "de", "es", "zh"])
def test_multicurrency_template_cells_are_translated_and_numeric(tmp_path: Path, language: str) -> None:
    from okama_planner.localization import translate
    from okama_planner.reports import export_report

    export_report([saved_multicurrency()], tmp_path / "plan.xlsx", language=language)
    book = load_workbook(tmp_path / "plan.xlsx")
    assert translate("Currency groups", language) in book.sheetnames
    assert f"{translate('Cash Flow', language)} USD" in book.sheetnames
    cash = book[translate("Cash Flow", language)]
    assert cash["A6"].value == translate("TOTAL INCOME", language)
    assert book[translate("Summary", language)]["B7"].value == 10100
    assert "[\u0024-" in book[translate("Summary", language)]["B7"].number_format
    assert book[translate("Goals", language)]["D5"].value == translate("Month", language)


@pytest.mark.parametrize(
    ("language", "locale"),
    [("en", "en-US"), ("ru", "ru-RU"), ("de", "de-DE"), ("es", "es-ES"), ("zh", "zh-CN")],
)
def test_html_selects_locale_without_changing_saved_chart_data(
    tmp_path: Path, language: str, locale: str
) -> None:
    from okama_planner.charts import export_charts
    from test_charts import payload, saved_result

    original = saved_result()
    document = export_charts(original, tmp_path, language=language)[0].read_text()
    assert payload(document)["locale"] == locale
    assert payload(document)["charts"] == original["charts"]
    assert "toLocaleString('en-US')" not in document


def test_export_errors_follow_selected_language_and_leave_no_artifact(tmp_path: Path) -> None:
    from okama_planner.charts import export_charts
    from okama_planner.reports import export_report
    from test_reports import scenario
    from test_charts import saved_result

    bad = saved_result()
    bad["currency"] = "INVALID"
    with pytest.raises(ValueError) as error:
        export_charts(bad, tmp_path / "charts", language="ru")
    assert "Не удалось" in str(error.value)
    item = scenario()
    item["result"]["currency"] = "EUR"
    with pytest.raises(ValueError) as error:
        export_report([item], tmp_path / "bad.xlsx", language="de")
    assert "Vorgang" in str(error.value)
    assert not (tmp_path / "bad.xlsx").exists()
    assert not (tmp_path / "charts").exists()


def test_invalid_saved_request_returns_localized_value_error(tmp_path: Path) -> None:
    from okama_planner.reports import export_report
    from test_reports import scenario

    item = scenario()
    item["request"]["currency"] = "PRIVATE_INVALID"
    with pytest.raises(ValueError) as error:
        export_report([item], tmp_path / "bad.xlsx", language="ru")
    assert "формату" in str(error.value)
    assert "PRIVATE_INVALID" not in str(error.value)
    assert not (tmp_path / "bad.xlsx").exists()


def test_portfolio_mode_display_is_translated_without_touching_request(tmp_path: Path) -> None:
    from okama_planner.reports import export_report
    from okama_planner.localization import translate
    from test_reports import scenario

    item = scenario()
    book = load_workbook(export_report([item], tmp_path / "mode.xlsx", language="de"))
    assert book[translate("Summary", "de")]["B8"].value == "Gemeinsames Portfolio"
    assert item["result"]["portfolio_mode"] == "single"
