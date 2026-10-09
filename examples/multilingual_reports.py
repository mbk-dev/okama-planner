"""Create synthetic national-currency editions using dated macro/history snapshots."""

import argparse
import json
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from okama_planner import forecast
from okama_planner.charts import export_charts
from okama_planner.localization import LANGUAGES, translate
from okama_planner.reports import ReportBrand, export_report

DEFAULT_CURRENCIES = {"en": "USD", "ru": "RUB", "zh": "CNY", "de": "EUR", "es": "EUR"}
SYNTHETIC_AMOUNTS = {
    "USD": (150000, 3000, 2300, 20000, 100000, 2000),
    "EUR": (130000, 3200, 2500, 18000, 90000, 2200),
    "CNY": (1100000, 22000, 16000, 130000, 700000, 14000),
    "RUB": (13000000, 250000, 180000, 1500000, 12000000, 150000),
}


def build_request(
    currency: str, snapshots: dict[str, Any], *, language: str = "en",
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Use independently authored fictional amounts and actual base-currency return samples."""
    metadata = deepcopy(snapshots["currencies"][currency])
    if metadata["risk_free_rate"] is None:
        raise ValueError(f"{currency} policy rate requires an explicit dated reference in the snapshot")
    request = json.loads((Path(__file__).parent / "baseline-request.json").read_text())
    opening, income, expenses, car, home, pension = SYNTHETIC_AMOUNTS[currency]
    metadata["amount_method"] = "Independently authored synthetic amounts; no FX conversion"
    request["currency"] = currency
    plan = request["plan"]
    plan.update(t0="2026-10", retirement_year=2040, withdrawal_years=30, horizon_years=44,
                accumulation_last_date_pin="2026-09", withdrawal_last_date_pin="2026-09")
    for key in plan["rates"]:
        plan["rates"][key] = (metadata["risk_free_rate"] if key == "buffer_rate"
                              else metadata["inflation_rate"])
    for asset in plan["assets"]:
        asset.update(amount=opening, currency=currency,
                     label=translate("Synthetic invested capital", language))
    plan["persons"] = [{"name": translate("Participant 1", language), "birth_year": 1997, "role": "adult"}]
    for item in plan["budget_items"]:
        item["monthly_amount"] = income if item["kind"] == "income" else expenses
        item["label"] = translate("Employment income" if item["kind"] == "income" else "Household expenses",
                                  language)
    plan["budget_items"].append({"kind": "income", "label": translate("State pension", language),
                                 "monthly_amount": expenses * .2, "start_month": "2040-10"})
    plan["goals"] = [
        {"goal_id": 1, "label": translate("Car", language), "kind": "lump", "amount_pv": car,
         "pv_year": 2026, "target_year": 2029, "target_month": 10, "becomes_asset": True},
        {"goal_id": 2, "label": translate("Home purchase", language), "kind": "lump", "amount_pv": home,
         "pv_year": 2026, "target_year": 2031, "target_month": 10, "becomes_asset": True},
        {"goal_id": 3, "label": translate("Early retirement", language), "kind": "retirement_income",
         "amount_pv": pension, "pv_year": 2026},
    ]
    request["return_samples"] = {
        stage: {"start_month": portfolio["start_month"], "monthly_returns": portfolio["monthly_returns"]}
        for stage, portfolio in metadata["portfolios"].items()
    }
    return request, metadata


def _metadata(path: Path, metadata: dict[str, Any], language: str) -> None:
    """Append reproducible economic assumptions without altering saved forecast values."""
    book = load_workbook(path)
    sheet = book[translate("Assumptions", language)]
    rows = [("Average inflation", metadata["inflation_rate"], True),
            ("Reference rate", metadata["risk_free_rate"], True)]
    rows.extend((key, json.dumps(value, ensure_ascii=False, allow_nan=False), False)
                for key, value in metadata.items()
                if key not in {"inflation_rate", "risk_free_rate", "portfolios"}
                and not key.startswith("fx_"))
    for stage, portfolio in metadata["portfolios"].items():
        for key, caption in (("annual_return", "Annual return"), ("annual_risk", "Annual risk"),
                             ("cagr", "Historical CAGR")):
            rows.append((f"{stage}: {translate(caption, language)}", portfolio[key], True))
        for key, value in portfolio.items():
            if key not in {"annual_return", "annual_risk", "cagr", "monthly_returns"}:
                rows.append((f"{stage}.{key}", json.dumps(value, ensure_ascii=False, allow_nan=False), False))
    for key, value, rate in rows:
        row = sheet.max_row + 1
        values = (translate("Economic data", language), translate(key, language), value)
        for column, text in enumerate(values, 1):
            cell = sheet.cell(row, column, text)
            cell.font = Font(name="Microsoft YaHei" if language == "zh" else "Arial", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(text, str):
                cell.data_type = "s"
        if rate:
            sheet.cell(row, 3).number_format = "0.00%"
        sheet.row_dimensions[row].height = 30
    sheet.print_area = sheet.dimensions
    book.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--languages", nargs="+", choices=LANGUAGES, default=list(LANGUAGES))
    parser.add_argument("--output-dir", type=Path, default=Path.cwd() / "tmp" / "multilingual-reports")
    args = parser.parse_args()
    snapshots = json.loads((Path(__file__).parent / "currency-assumptions.json").read_text())
    requests = {
        language: build_request(DEFAULT_CURRENCIES[language], snapshots, language=language)
        for language in args.languages
    }
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for language, (request, metadata) in requests.items():
        currency = request["currency"]
        result = forecast(request)
        for kind, value in (("request", request), ("result", result)):
            snapshot = args.output_dir / f"{language}-{currency.lower()}-{kind}.json"
            snapshot.write_text(json.dumps(value, indent=2, allow_nan=False) + "\n")
        images = {}
        for logarithmic in (False, True):
            scale = "log" if logarithmic else "linear"
            paths = export_charts(result, args.output_dir / language / scale,
                                  format="png", logarithmic=logarithmic, language=language)
            images.update({path.stem + ("_log" if logarithmic else ""): path for path in paths})
        scenarios = [{"label": translate("Baseline", language), "request": request, "result": result}]
        path = export_report(scenarios,
                             args.output_dir / f"plan-{language}.xlsx", language=language,
                             chart_images=images, brand=ReportBrand(contact=""))
        _metadata(path, metadata, language)
    print(f"Synthetic language editions written to {args.output_dir}")


if __name__ == "__main__":
    main()
